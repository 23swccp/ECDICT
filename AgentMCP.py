"""性能强悍的MCP脚本,给agent开子代理批量处理词条文件提供可行工具

---------------------------------------------------------------------------------------------------------------
---------------------------------------------------------------------------------------------------------------

实测3个 GPT 6.0 Luna 标注150词的中文:word以 CEFR B2—C1 为主，夹有学术性偏难的词

吞吐 53.6词/分钟，折算 3,215词/小时,包含联网搜索


需要
    python -m pip install "mcp>=2.3,<3" "openpyxl>=3.1.5,<4"

只启动一个共享服务，所有代理都连接 http://127.0.0.1:8765/mcp, 通过MCP Streamable HTTP 协议调用 batch 工具

操作对象是一个.xlsx文件,一般只需要两个字段:(word,你想修改的字段如translation)，这样比较好操作,当然多个字段也可以.

---------------------------------------------------------------------------------------------------------------
---------------------------------------------------------------------------------------------------------------

比如你发现数据库中的单词:(apple, translate, beautiful, quickly)缺失(pos, translation):
1.整理成.xlsx文件:
| word      | pos | translation |
|-----------|-----|-------------|
| apple     |     |             |
| translate |     |             |
| beautiful |     |             |
| quickly   |     |             |

第一行必须是表头，不能有空列或重复列。

2.启动服务，每批分配两行： python AgentMCP.py 待标注示例.xlsx --batch-size 2


3.在支持 MCP 的宿主中连接： http://127.0.0.1:8765/mcp


4.给主代理下达任务：
    安排两个子代理，通过同一个 AgentMCP 服务领取批次，
    补充 pos 和 translation。pos 使用 n.、v.、adj.、adv. 等缩写；

主代理可以用 action=status 查看进度，或者用 action=release 归还未提交的批次。


5.返回
| word      | pos  | translation |
|-----------|------|-------------|
| apple     | n.   | 苹果        |
| translate | v.   | 翻译        |
| beautiful | adj. | 美丽的      |
| quickly   | adv. | 迅速地      |

"""




from __future__ import annotations


DEFAULT_BATCH_SIZE = 50                 # 每批行数，1～1000
DEFAULT_MAX_PENDING = 128               # 保存队列容量，1～10000 批
DEFAULT_MERGE_MS = 50                   # 合并保存等待时间，0～1000 毫秒
DEFAULT_SHEET = None                    # 工作表名称，None 表示第一个
DEFAULT_PORT = 8765                     # HTTP 端口；修改后同步更新连接地址
DEFAULT_TRANSPORT = "streamable-http"    # 连接方式：streamable-http 或 stdio

# 主 agent：先试跑，再按耗时、质量和保存积压调整子代理数量，避免盲目增加并发。
# 用 batch(action="status", request_id="唯一编号") 查询进度，避免频繁轮询；done=true 才算完成。
# 代理退出后，用 action="release" 和 batch_id 归还未提交批次；保存失败以原请求、原内容重试。
# 每批行数由启动参数 --batch-size 设置，不支持运行中修改；不要拆批提交或中途重启。
# 提交成功会自动领取下一批；代理结束前须归还新领批次，主 agent 完成后抽查质量。

import argparse
import asyncio
from collections import deque
from contextlib import asynccontextmanager
from copy import deepcopy
from dataclasses import dataclass, field
import hashlib
import json
import math
import os
from pathlib import Path
import socket
import sys
import tempfile
from typing import Literal
import uuid


CellValue = str | int | float | bool | None


class BatchError(ValueError):
    """请求被拒绝或批次失效。"""


class SaveError(BatchError):
    """提交仍在内存，须以原请求、原内容重试。"""


def canonical(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True,
                      separators=(",", ":"), allow_nan=False)


def validate_cell(value, location):
    if value is not None and type(value) not in (str, int, float, bool):
        raise BatchError(f"{location}: 仅支持文本、数字、布尔值或 null。")
    if type(value) is float and not math.isfinite(value):
        raise BatchError(f"{location}: 不支持 NaN 或 Infinity。")
    if isinstance(value, str):
        if len(value) > 32767:
            raise BatchError(f"{location}: 文本超过 32767 字符。")
        for char in value:
            code = ord(char)
            if (code < 32 and char not in "\t\n\r") or 0xD800 <= code <= 0xDFFF or code in (0xFFFE, 0xFFFF):
                raise BatchError(f"{location}: 含非法 XML 字符。")


def fingerprint(path):
    """检测外部改写，包括文件大小未变的修改。"""
    with path.open("rb") as handle:
        before = os.fstat(handle.fileno())
        digest = hashlib.file_digest(handle, "sha256").hexdigest()
        after = os.fstat(handle.fileno())
    identity = lambda stat: (stat.st_dev, stat.st_ino, stat.st_size, stat.st_mtime_ns)
    if identity(before) != identity(after) or identity(after) != identity(path.stat()):
        raise BatchError("Excel 正被外部修改；停止编辑后重启服务。")
    return (*identity(after), digest)


class WorkbookOwner:
    """按工作簿路径占用本机端口，实现跨进程独占；端口冲突时拒绝启动。"""

    def __init__(self, path):
        identity = os.path.normcase(str(path.resolve())).encode("utf-8")
        port = 40000 + int.from_bytes(hashlib.sha256(identity).digest()[:4], "big") % 20000
        self.socket = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
        try:
            if hasattr(socket, "SO_EXCLUSIVEADDRUSE"):
                self.socket.setsockopt(socket.SOL_SOCKET, socket.SO_EXCLUSIVEADDRUSE, 1)
            self.socket.bind(("127.0.0.1", port))
        except OSError as exc:
            self.close()
            raise BatchError(f"工作簿被占用或锁端口 {port} 不可用。") from exc

    def close(self):
        self.socket.close()


@dataclass
class Request:
    signature: str
    future: asyncio.Future
    response: dict | None = None
    submission: Claim | None = None


@dataclass
class Claim:
    batch_id: str
    indices: tuple[int, ...]
    original: list[dict]
    state: str = "claimed"
    submitted: list[dict] | None = None
    submitted_key: str | None = None
    request: Request | None = None
    changes: list[tuple[int, int, CellValue]] = field(default_factory=list)


class BatchStore:
    """单个事件循环调度批次，仅写入线程修改单元格。"""

    def __init__(self, path, sheet=DEFAULT_SHEET, batch_size=DEFAULT_BATCH_SIZE,
                 max_pending=DEFAULT_MAX_PENDING, merge_ms=DEFAULT_MERGE_MS):
        if type(batch_size) is not int or not 1 <= batch_size <= 1000:
            raise BatchError("batch-size 须为 1～1000 的整数。")
        if type(max_pending) is not int or not 1 <= max_pending <= 10000:
            raise BatchError("max-pending 须为 1～10000 的整数。")
        if not 0 <= merge_ms <= 1000:
            raise BatchError("merge-ms 须在 0～1000 之间。")
        self.path = Path(path).resolve(strict=True)
        if self.path.suffix.lower() != ".xlsx":
            raise BatchError("输入必须是 .xlsx 文件。")
        self.owner = WorkbookOwner(self.path)
        self.workbook = None
        try:
            from openpyxl import load_workbook
            self.expected = fingerprint(self.path)
            self.workbook = load_workbook(self.path, data_only=False, keep_links=True, rich_text=True)
            if fingerprint(self.path) != self.expected:
                raise BatchError("读取时 Excel 已变化，请重启服务。")
            if sheet is not None and sheet not in self.workbook.sheetnames:
                raise BatchError(f"工作表不存在：{sheet}")
            self.sheet = self.workbook[sheet] if sheet else self.workbook.worksheets[0]
            if self.sheet.merged_cells.ranges:
                raise BatchError("不支持合并单元格。")
            self.headers = [self.sheet.cell(1, col).value for col in range(1, self.sheet.max_column + 1)]
            if any(not isinstance(h, str) or not h.strip() for h in self.headers):
                raise BatchError("首行须为非空文本表头。")
            if len(set(self.headers)) != len(self.headers):
                raise BatchError("表头不能重复。")
            for header in self.headers:
                validate_cell(header, "表头")
            self.data = []
            for cells in self.sheet.iter_rows(min_row=2, max_col=len(self.headers)):
                values = []
                for cell in cells:
                    if cell.data_type in ("f", "e"):
                        raise BatchError(f"{cell.coordinate}: 不支持公式或 Excel 错误值。")
                    validate_cell(cell.value, cell.coordinate)
                    values.append(cell.value)
                if any(v is not None and v != "" for v in values):
                    self.data.append((cells[0].row, dict(zip(self.headers, values))))
        except BaseException:
            if self.workbook is not None:
                self.workbook.close()
            self.owner.close()
            raise
        self.batch_size = batch_size
        self.merge_seconds = merge_ms / 1000
        self.todo = deque(range(len(self.data)))
        self.claims: dict[str, Claim] = {}
        self.requests: dict[str, Request] = {}
        self.queue: asyncio.Queue[Claim] = asyncio.Queue(maxsize=max_pending)
        self.lock = asyncio.Lock()
        self.writer = None
        self.completed = 0
        self.save_count = 0
        self.last_save_error = None
        self.closing = False
        self.closed = False

    @staticmethod
    def _future():
        future = asyncio.get_running_loop().create_future()
        # 客户端断开后仍可重试，并避免异常无人读取的警告。
        future.add_done_callback(lambda f: f.exception() if not f.cancelled() else None)
        return future

    def _claim(self):
        if self.closing:
            return {"waiting": True}
        if not self.todo:
            return {"done": True} if self.completed == len(self.data) else {"waiting": True}
        indices = tuple(self.todo.popleft() for _ in range(min(self.batch_size, len(self.todo))))
        original = [deepcopy(self.data[i][1]) for i in indices]
        batch_id = uuid.uuid4().hex
        self.claims[batch_id] = Claim(batch_id, indices, original)
        return {"batch_id": batch_id, "rows": deepcopy(original)}

    def status(self):
        active = [c for c in self.claims.values() if c.state in ("claimed", "queued", "saving", "save_failed")]
        return {
            "total": len(self.data), "pending": len(self.todo),
            "processing": sum(len(c.indices) for c in active if c.state == "claimed"),
            "pending_save": sum(len(c.indices) for c in active if c.state != "claimed"),
            "saved": self.completed, "save_count": self.save_count,
            "active_batches": [{"batch_id": c.batch_id, "state": c.state, "rows": len(c.indices)} for c in active],
            "last_save_error": self.last_save_error,
            "done": self.completed == len(self.data),
        }

    def _validate_rows(self, claim, rows):
        if not isinstance(rows, list) or len(rows) != len(claim.indices):
            raise BatchError(f"rows 须完整保留本批 {len(claim.indices)} 行及顺序。")
        expected = set(self.headers)
        for number, row in enumerate(rows, 1):
            if not isinstance(row, dict) or set(row) != expected:
                raise BatchError(f"第 {number} 行字段须与表头一致。")
            for name, value in row.items():
                validate_cell(value, f"本批第 {number} 行/{name}")

    def _enqueue(self, claim):
        if self.queue.full():
            raise BatchError("保存队列已满；稍后以原请求、原内容重试。")
        claim.state = "queued"
        self.queue.put_nowait(claim)
        if self.writer is None or self.writer.done():
            self.writer = asyncio.create_task(self._write_loop(), name="AgentMCP Excel writer")

    def _replay(self, record):
        response = record.response
        if response and "batch_id" in response and self.claims[response["batch_id"]].state == "released":
            raise BatchError("原批次已归还；用新 request_id 领取。")
        if record.submission is not None and record.submission.state == "save_failed":
            # 多个请求共享此记录，重试只入队一次。
            if self.queue.full():
                raise BatchError("保存队列已满；稍后以原请求、原内容重试。")
            record.future = self._future()
            self._enqueue(record.submission)
        return record.future

    async def batch(self, request_id, batch_id=None, rows=None, action="next"):
        if not isinstance(request_id, str) or not request_id.strip() or len(request_id) > 256:
            raise BatchError("request_id 须为 1～256 字符的非空字符串。")
        if action not in ("next", "status", "release"):
            raise BatchError("action 须为 next、status 或 release。")
        if batch_id is not None and (not isinstance(batch_id, str) or not batch_id):
            raise BatchError("batch_id 须为非空字符串。")
        if action == "status" and (batch_id is not None or rows is not None):
            raise BatchError("status 不接受 batch_id 或 rows。")
        if action == "release" and (batch_id is None or rows is not None):
            raise BatchError("release 须带 batch_id，不带 rows。")
        if action == "next" and ((batch_id is None) != (rows is None)):
            raise BatchError("batch_id、rows：领取时均省略，提交时均提供。")
        try:
            signature = canonical({"action": action, "batch_id": batch_id, "rows": rows})
        except (TypeError, ValueError) as exc:
            raise BatchError("参数须为合法 JSON，不含 NaN 或 Infinity。") from exc
        async with self.lock:
            if self.closing or self.closed:
                raise BatchError("服务关闭中或已关闭。")
            if request_id in self.requests:
                record = self.requests[request_id]
                if record.signature != signature:
                    raise BatchError("request_id 冲突：重试须保持原参数。")
                future = self._replay(record)
            elif action == "status":
                # 每次查询返回最新进度。
                return deepcopy(self.status())
            elif action == "release":
                claim = self.claims.get(batch_id)
                if claim is None or claim.state not in ("claimed", "released"):
                    raise BatchError("批次不存在或已提交，不能归还。")
                if claim.state == "claimed":
                    claim.state = "released"
                    self.todo.extendleft(reversed(claim.indices))
                record = Request(signature, self._future(), response={"released": True})
                record.future.set_result(record.response)
                self.requests[request_id] = record
                future = record.future
            elif batch_id is None:
                record = Request(signature, self._future(), response=self._claim())
                record.future.set_result(record.response)
                self.requests[request_id] = record
                future = record.future
            else:
                claim = self.claims.get(batch_id)
                if claim is None or claim.state == "released":
                    raise BatchError("批次已失效；用新 request_id 重新领取。")
                self._validate_rows(claim, rows)
                submitted_key = canonical(rows)
                if claim.request is not None:
                    if claim.submitted_key != submitted_key:
                        raise BatchError("批次内容冲突；请重试原提交。")
                    record = claim.request
                    future = self._replay(record)
                    self.requests[request_id] = record
                else:
                    if self.queue.full():
                        raise BatchError("保存队列已满；稍后以原请求、原内容重试。")
                    record = Request(signature, self._future(), submission=claim)
                    claim.submitted = deepcopy(rows)
                    claim.submitted_key = submitted_key
                    claim.changes = [
                        (self.data[index][0], col, row[name])
                        for index, old, row in zip(claim.indices, claim.original, claim.submitted)
                        for col, name in enumerate(self.headers, 1)
                        if canonical(old[name]) != canonical(row[name])
                    ]
                    claim.request = record
                    self.requests[request_id] = record
                    self._enqueue(claim)
                    future = record.future
        # 调用方断开或取消时，保留已接受的任务和重试共享的 Future。
        result = await asyncio.shield(future)
        return deepcopy(result)

    def _check_external(self):
        if fingerprint(self.path) != self.expected:
            raise OSError("Excel 已被外部修改，拒绝覆盖；停止服务并检查文件。")

    def _save_group(self, claims):
        """在唯一写入线程保存，失败时回滚内存。"""
        self._check_external()
        undo = []
        temporary = None
        try:
            for claim in claims:
                for row, col, value in claim.changes:
                    cell = self.sheet.cell(row, col)
                    undo.append((cell, cell.value, cell.data_type))
                    cell.value = value
                    if isinstance(value, str):
                        # 以文本保存，避免被识别为公式或错误值。
                        cell.data_type = "s"
            if undo:
                mode = self.path.stat().st_mode
                descriptor, temporary = tempfile.mkstemp(prefix=f".{self.path.stem}.AgentMCP-", suffix=".xlsx", dir=self.path.parent)
                os.close(descriptor)
                self.workbook.save(temporary)
                os.chmod(temporary, mode)
                with open(temporary, "r+b") as handle:
                    os.fsync(handle.fileno())
                new_fingerprint = fingerprint(Path(temporary))
                self._check_external()
                os.replace(temporary, self.path)
                temporary = None
                # 替换后沿用临时文件指纹，确认成功前不再访问磁盘。
                self.expected = new_fingerprint
        except BaseException:
            for cell, value, data_type in reversed(undo):
                cell.value = value
                cell.data_type = data_type
            raise
        finally:
            if temporary is not None:
                try:
                    os.unlink(temporary)
                except OSError:
                    pass
        return bool(undo)

    async def _write_loop(self):
        while True:
            first = await self.queue.get()
            if self.merge_seconds:
                await asyncio.sleep(self.merge_seconds)
            claims = [first]
            while not self.queue.empty():
                claims.append(self.queue.get_nowait())
            async with self.lock:
                for claim in claims:
                    claim.state = "saving"
            try:
                changed = await asyncio.to_thread(self._save_group, claims)
            except Exception as exc:
                async with self.lock:
                    self.last_save_error = str(exc)
                    for claim in claims:
                        claim.state = "save_failed"
                        claim.request.future.set_exception(SaveError(
                            f"保存失败，结果保留在内存；以原 request_id、原内容重试：{exc}"
                        ))
            else:
                async with self.lock:
                    self.save_count += int(changed)
                    # 整组标记完成后再分配，确保 waiting/done 状态准确。
                    for claim in claims:
                        claim.state = "saved"
                        self.completed += len(claim.indices)
                    if not any(c.state == "save_failed" for c in self.claims.values()):
                        self.last_save_error = None
                    for claim in claims:
                        record = claim.request
                        record.response = self._claim()
                        record.future.set_result(record.response)
            finally:
                for _ in claims:
                    self.queue.task_done()

    async def aclose(self):
        if self.closed:
            return
        async with self.lock:
            self.closing = True
        await self.queue.join()
        if self.writer is not None:
            self.writer.cancel()
            try:
                await self.writer
            except asyncio.CancelledError:
                pass
        failed = sum(len(c.indices) for c in self.claims.values() if c.state == "save_failed")
        if failed:
            print(f"关闭时仍有 {failed} 行未保存：{self.last_save_error}", file=sys.stderr)
        self.workbook.close()
        self.owner.close()
        self.closed = True


def create_server(store):
    try:
        from mcp.server import MCPServer
        from mcp.server.mcpserver.exceptions import ToolError
    except ImportError as exc:
        raise RuntimeError('需要 MCP SDK 2.x：python -m pip install "mcp>=2.3,<3"') from exc

    @asynccontextmanager
    async def lifespan(server):
        try:
            yield {}
        finally:
            await store.aclose()

    server = MCPServer("AgentMCP Excel annotation", lifespan=lifespan)

    @server.tool(structured_output=True)
    async def batch(request_id: str, batch_id: str | None = None,
                    rows: list[dict[str, CellValue]] | None = None,
                    action: Literal["next", "status", "release"] = "next") -> dict[str, object]:
        
        try:
            return await store.batch(request_id, batch_id, rows, action)
        except (BatchError, OSError) as exc:
            raise ToolError(str(exc)) from exc

    return server


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("excel", type=Path, help="输入暨待审批输出 .xlsx")
    parser.add_argument("--sheet", default=DEFAULT_SHEET, help="工作表名称，未设置时用第一个")
    parser.add_argument("--batch-size", type=int, default=DEFAULT_BATCH_SIZE, help="每批行数（默认：%(default)s）")
    parser.add_argument("--max-pending", type=int, default=DEFAULT_MAX_PENDING, help="待保存队列容量（默认：%(default)s）")
    parser.add_argument("--merge-ms", type=int, default=DEFAULT_MERGE_MS, help="合并保存等待窗口，0～1000 毫秒（默认：%(default)s）")
    parser.add_argument("--transport", choices=("streamable-http", "stdio"), default=DEFAULT_TRANSPORT, help="连接方式（默认：%(default)s）")
    parser.add_argument("--port", type=int, default=DEFAULT_PORT, help="HTTP 端口，仅监听 127.0.0.1（默认：%(default)s）")
    parser.add_argument("--check", action="store_true", help="只验证并显示表格信息，不修改文件")
    args = parser.parse_args()
    if not 1 <= args.port <= 65535:
        parser.error("port 必须在 1 到 65535 之间。")
    store = None
    try:
        store = BatchStore(args.excel, args.sheet, args.batch_size, args.max_pending, args.merge_ms)
        if args.check:
            print(canonical({"file": str(store.path), "sheet": store.sheet.title,
                             "headers": store.headers, "rows": len(store.data), "batch_size": store.batch_size}))
            return 0
        server = create_server(store)
        print(f"AgentMCP：{store.path}，{len(store.data)} 行，每批最多 {store.batch_size} 行。", file=sys.stderr)
        print("状态仅在内存；原表回填，结果待审批；没有自动领取超时。", file=sys.stderr)
        if args.transport == "streamable-http":
            print(f"所有代理连接同一个地址：http://127.0.0.1:{args.port}/mcp", file=sys.stderr)
            server.run(transport="streamable-http", host="127.0.0.1", port=args.port)
        else:
            server.run(transport="stdio")
        return 0
    except (BatchError, OSError, RuntimeError, ImportError) as exc:
        print(f"AgentMCP 错误：{exc}", file=sys.stderr)
        return 1
    finally:
        if store is not None and not store.closed:
            # 服务退出时会排空队列；此处清理仅检查或启动失败时的资源。
            store.workbook.close()
            store.owner.close()


if __name__ == "__main__":
    for stream in (sys.stdin, sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8")
    raise SystemExit(main())


# ====================================================================
# 启动参数示例（默认值见文件开头）
#
# --batch-size 50      每批最多行数，范围1～1000。
# --sheet 词条         指定工作表，默认第一个。
# --port 8765          HTTP端口，范围1～65535。
# 修改后同步调整MCP连接地址。
# --merge-ms 50        合并保存等待窗口，范围0～1000毫秒。
# 
# --max-pending 128    待保存队列容量，范围1～10000批。
# 
# --transport streamable-http
# 默认HTTP模式，也可选择stdio。
# stdio要求代理共享同一个服务进程。
# --check              只检查表格，不修改文件。
#
# batch工具的调用参数
#
# request_id   必填，1～256字符，同一服务中保持唯一。
# 新操作用新编号；通信重试用原编号和原内容。
# batch_id     服务返回的批次编号，提交或归还时原样使用。
# rows         提交本批完整数据，保留字段、行数和行顺序。
# null表示清空单元格
# action       next(默认)、status、release。
#
# 领取: {"request_id": "agent1-001"}
# 提交: 新request_id + 原batch_id + 完整rows。
# 先保存本批，成功后自动返回下一批。
# 查询: {"request_id": "status-001", "action": "status"}
# 归还: 新request_id + action="release" + 原batch_id。
# 只归还尚未提交的批次
#
# 领取或提交后的回复
#
# batch_id、rows    已分配一批，继续处理。
# waiting=true      暂时无批次可领，其他批次仍在处理。
# 稍后使用新的request_id再次领取。
# done=true         全部批次已完成，所需修改已保存。
#

# ====================================================================
